"""Trích xuất văn bản từ tài liệu: PDF, DOCX, XLSX và ảnh.

Nguyên tắc suy giảm dần: luôn cố lấy được phần văn bản lớp chữ trước; chỉ gọi OCR
cho phần còn thiếu. Nếu máy chưa cài thư viện hoặc chưa có Tesseract thì vẫn trả
kết quả kèm cảnh báo thay vì làm hỏng cả tài liệu.
"""

from __future__ import annotations

import logging
import shutil
import subprocess
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from ..config import get_settings

logger = logging.getLogger("mindora.extraction")

# Dưới ngưỡng này coi như trang là ảnh scan, cần dùng OCR.
MIN_CHARS_PER_PAGE = 40
_TESSERACT_CANDIDATES = (
    r"D:\TroLiThongMinh\tools\tesseract\tesseract.exe",
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
)


class ExtractionError(Exception):
    """Không trích xuất được văn bản."""


class MissingDependency(ExtractionError):
    """Thiếu thư viện cần thiết để đọc loại tệp này."""


@dataclass(frozen=True)
class ExtractionResult:
    pages: list[tuple[int | None, str]]
    page_count: int
    method: str
    warnings: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(page_text for _, page_text in self.pages if page_text.strip())

    @property
    def is_empty(self) -> bool:
        return not self.text.strip()


@lru_cache(maxsize=1)
def find_tesseract() -> str | None:
    """Tìm tesseract.exe: biến cấu hình, PATH, rồi các vị trí cài đặt thông dụng."""
    settings = get_settings()
    if settings.tesseract_cmd:
        candidate = Path(settings.tesseract_cmd)
        return str(candidate) if candidate.exists() else None
    found = shutil.which("tesseract")
    if found:
        return found
    for candidate in _TESSERACT_CANDIDATES:
        if Path(candidate).exists():
            return candidate
    return None


@lru_cache(maxsize=1)
def ocr_languages() -> frozenset[str]:
    """Danh sách ngôn ngữ Tesseract đang có, dùng để cảnh báo khi thiếu gói tiếng Việt."""
    command = find_tesseract()
    if not command:
        return frozenset()
    try:
        result = subprocess.run(  # noqa: S603 - lệnh cục bộ do cấu hình quy định
            [command, "--list-langs"], capture_output=True, text=True, timeout=15, check=False
        )
    except (OSError, subprocess.SubprocessError):
        return frozenset()
    return frozenset(line.strip() for line in result.stdout.splitlines()[1:] if line.strip())


def configured_languages() -> list[str]:
    """Danh sách ngôn ngữ theo cấu hình, ví dụ "vie+eng" -> ["vie", "eng"]."""
    return [code for code in get_settings().tesseract_lang.split("+") if code]


@lru_cache(maxsize=1)
def effective_languages() -> tuple[str, ...]:
    """Ngôn ngữ OCR thực sự dùng được: chỉ lấy những gói đã cài trên máy.

    Nhờ vậy khi thiếu gói tiếng Việt, hệ thống vẫn OCR được phần chữ Latinh
    thay vì báo lỗi cứng và làm hỏng cả tài liệu.
    """
    installed = ocr_languages()
    usable = [code for code in configured_languages() if code in installed]
    if usable:
        return tuple(usable)
    if "eng" in installed:
        return ("eng",)
    return ()


def ocr_status() -> dict:
    """Thông tin OCR dùng cho endpoint kiểm tra tình trạng."""
    command = find_tesseract()
    installed = sorted(ocr_languages())
    required = configured_languages()
    missing = [code for code in required if code not in installed]
    return {
        "available": command is not None,
        "command": command,
        "languages": installed,
        "required_languages": required,
        "effective_languages": list(effective_languages()),
        "missing_languages": missing,
    }


def _ocr_image_file(image_path: Path) -> str:
    """Chạy OCR trên một tệp ảnh, trả về văn bản (có thể rỗng)."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:  # pragma: no cover - phụ thuộc môi trường
        raise MissingDependency("Thiếu pytesseract hoặc Pillow để chạy OCR.") from exc
    command = find_tesseract()
    if command:
        pytesseract.pytesseract.tesseract_cmd = command
    languages = effective_languages()
    if not languages:
        raise MissingDependency(
            "Tesseract chưa có gói ngôn ngữ nào (cần ít nhất eng hoặc vie)."
        )
    with Image.open(image_path) as image:
        return pytesseract.image_to_string(image, lang="+".join(languages))


def _ocr_warnings(status: dict) -> list[str]:
    """Cảnh báo dễ hiểu cho người dùng khi chất lượng OCR có thể bị giới hạn."""
    warnings: list[str] = []
    if status["missing_languages"]:
        warnings.append(
            "Thiếu gói ngôn ngữ Tesseract: " + ", ".join(status["missing_languages"]) + "."
        )
    required = status["required_languages"]
    effective = status["effective_languages"]
    if effective and "vie" in required and "vie" not in effective:
        warnings.append("Chưa có gói tiếng Việt nên chất lượng OCR tiếng Việt sẽ thấp.")
    return warnings


def _pdf_text_pages(path: Path) -> list[tuple[int, str]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise MissingDependency("Thiếu thư viện pypdf để đọc tệp PDF.") from exc
    try:
        reader = PdfReader(str(path))
        pages: list[tuple[int, str]] = []
        for index, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:  # noqa: BLE001 - một trang hỏng không nên làm hỏng cả tệp
                logger.warning("Không đọc được lớp chữ trang %d của %s", index, path.name)
                text = ""
            pages.append((index, text))
        return pages
    except MissingDependency:
        raise
    except Exception as exc:
        raise ExtractionError(f"Không đọc được tệp PDF: {exc}") from exc


def _ocr_pdf_page(path: Path, page_number: int) -> str:
    """Dựng trang PDF thành ảnh rồi OCR, dùng cho trang bản scan."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return ""
    temp_dir = get_settings().upload_tmp_dir
    temp_dir.mkdir(parents=True, exist_ok=True)
    image_path = temp_dir / f"ocr-{page_number}-{abs(hash(str(path)))}.png"
    try:
        with fitz.open(str(path)) as document:
            page = document[page_number - 1]
            pixmap = page.get_pixmap(dpi=get_settings().ocr_dpi)
            pixmap.save(str(image_path))
        return _ocr_image_file(image_path)
    except Exception as exc:  # noqa: BLE001 - OCR là bước bổ sung, lỗi không được chặn tài liệu
        logger.warning("OCR trang %d thất bại (%s): %s", page_number, path.name, exc)
        return ""
    finally:
        image_path.unlink(missing_ok=True)


def extract_pdf(path: Path) -> ExtractionResult:
    """Lấy lớp chữ của PDF; trang nào quá nghèo thì bổ sung bằng OCR."""
    pages = _pdf_text_pages(path)
    warnings: list[str] = []
    method = "text"
    status = ocr_status()
    enriched: list[tuple[int | None, str]] = []

    for page_number, text in pages:
        if len(text.strip()) >= MIN_CHARS_PER_PAGE:
            enriched.append((page_number, text))
            continue
        if not status["available"]:
            enriched.append((page_number, text))
            continue
        ocr_text = _ocr_pdf_page(path, page_number)
        if ocr_text.strip():
            method = "text+ocr"
            enriched.append((page_number, ocr_text))
        else:
            enriched.append((page_number, text))

    if not status["available"]:
        warnings.append("Chưa cài Tesseract nên trang dạng scan không đọc được.")
    else:
        warnings.extend(_ocr_warnings(status))
    return ExtractionResult(pages=enriched, page_count=len(pages), method=method, warnings=warnings)


def extract_docx(path: Path) -> ExtractionResult:
    try:
        import docx
    except ImportError as exc:
        raise MissingDependency("Thiếu thư viện python-docx để đọc tệp DOCX.") from exc
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"Không đọc được tệp DOCX: {exc}") from exc
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return ExtractionResult(
        pages=[(None, "\n".join(lines))],
        page_count=1,
        method="text",
    )


def extract_xlsx(path: Path) -> ExtractionResult:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise MissingDependency("Thiếu thư viện openpyxl để đọc tệp XLSX.") from exc
    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError(f"Không đọc được tệp XLSX: {exc}") from exc
    pages: list[tuple[int | None, str]] = []
    try:
        for index, sheet in enumerate(workbook.worksheets, start=1):
            lines = [f"# {sheet.title}"]
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value).strip() for value in row]
                if any(cells):
                    lines.append(" | ".join(cells))
            pages.append((index, "\n".join(lines)))
    finally:
        workbook.close()
    return ExtractionResult(pages=pages, page_count=len(pages), method="text")


def extract_image(path: Path) -> ExtractionResult:
    status = ocr_status()
    if not status["available"]:
        raise MissingDependency("Chưa cài Tesseract nên chưa đọc được ảnh tài liệu.")
    try:
        text = _ocr_image_file(path)
    except MissingDependency:
        raise
    except Exception as exc:
        raise ExtractionError(f"Không đọc được ảnh: {exc}") from exc
    warnings = _ocr_warnings(status)
    return ExtractionResult(
        pages=[(1, text)],
        page_count=1,
        method="ocr",
        warnings=warnings,
    )


def extract(path: Path, file_type: str) -> ExtractionResult:
    """Điểm vào duy nhất: trả về văn bản theo từng trang."""
    if not path.exists():
        raise ExtractionError("Không tìm thấy tệp tài liệu trên đĩa.")
    handlers = {
        "pdf": extract_pdf,
        "docx": extract_docx,
        "xlsx": extract_xlsx,
        "image": extract_image,
    }
    handler = handlers.get(file_type)
    if not handler:
        raise ExtractionError(f"Không hỗ trợ trích xuất cho loại tệp {file_type}.")
    return handler(path)
