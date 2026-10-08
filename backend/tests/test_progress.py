"""Kiểm thử Module F: tiến độ học tập, hoạt động và gợi ý (dựa trên hỏi đáp)."""

from __future__ import annotations

from fastapi.testclient import TestClient

from .conftest import auth_header
from .test_chat import FakeProvider, use_provider
from .test_rag import _upload_knowledge


def test_progress_requires_authentication(client: TestClient) -> None:
    assert client.get("/api/progress").status_code == 401


def test_progress_is_empty_for_new_user(client: TestClient, make_user) -> None:
    token, _ = make_user()
    body = client.get("/api/progress", headers=auth_header(token)).json()

    assert body["overview"] == {
        "documents": 0,
        "questions": 0,
        "conversations": 0,
    }
    assert body["subjects"] == []
    assert body["suggestion"] is None
    assert body["activities"] == []
    assert body["daily_activity"] == []


def test_progress_aggregates_chat_and_documents(
    client: TestClient, monkeypatch, make_user, docx_factory
) -> None:
    token, _ = make_user()

    # Một câu hỏi chat.
    use_provider(monkeypatch, FakeProvider())
    client.post(
        "/api/chat",
        json={"message": "Giải thích gradient descent"},
        headers=auth_header(token),
    )

    # Một tài liệu đã sẵn sàng (chưa hỏi nên phải có gợi ý).
    _upload_knowledge(client, token, docx_factory)

    body = client.get("/api/progress", headers=auth_header(token)).json()
    overview = body["overview"]
    assert overview["documents"] == 1
    assert overview["questions"] == 1
    assert overview["conversations"] == 1

    # Lịch sử gộp hỏi đáp và tài liệu (FR-F1).
    assert {"Chat AI", "Tài liệu"} <= {item["type"] for item in body["activities"]}
    # Biểu đồ theo ngày có dữ liệu (FR-F3).
    assert body["daily_activity"]
    assert body["daily_activity"][0]["day"]
    # Câu hỏi chưa gắn môn nào nên nằm trong nhóm "Chung".
    assert any(subject["subject"] == "Chung" for subject in body["subjects"])


def test_suggestion_names_ready_document_without_questions(
    client: TestClient, make_user, upload
) -> None:
    """Tài liệu có subject_tag nhưng chưa hỏi câu nào → có gợi ý."""
    token, _ = make_user()
    import io

    import docx

    document = docx.Document()
    document.add_heading("Bài giảng Xác suất", level=1)
    for index in range(5):
        document.add_paragraph(
            f"Muc {index}. Phan phoi xac suat thong ke muc {index} rat quan trong."
        )
    buffer = io.BytesIO()
    document.save(buffer)

    response = upload(
        token,
        buffer.getvalue(),
        filename="xac-suat.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        subject_tag="Xác suất thống kê",
    )
    assert response.status_code == 201, response.text
    document_id = response.json()["id"]
    status = client.get(f"/api/documents/{document_id}/status", headers=auth_header(token)).json()
    assert status["status"] == "ready", status

    body = client.get("/api/progress", headers=auth_header(token)).json()
    assert body["suggestion"] is not None
    assert body["suggestion"]["subject"] == "Xác suất thống kê"


def test_suggestion_disappears_after_first_question(
    client: TestClient, monkeypatch, make_user, upload
) -> None:
    """Hỏi đủ một câu cho môn thì gợi ý không còn."""
    token, _ = make_user()
    use_provider(monkeypatch, FakeProvider())

    import io

    import docx

    document = docx.Document()
    document.add_heading("Bai giang Xac suat", level=1)
    for index in range(5):
        document.add_paragraph(f"Muc {index}. Phan phoi xac suat thong ke muc {index}.")
    buffer = io.BytesIO()
    document.save(buffer)
    response = upload(
        token,
        buffer.getvalue(),
        filename="xac-suat.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        subject_tag="Xác suất thống kê",
    )
    document_id = response.json()["id"]
    assert client.get(f"/api/documents/{document_id}/status", headers=auth_header(token)).json()["status"] == "ready"

    assert client.get("/api/progress", headers=auth_header(token)).json()["suggestion"] is not None

    # Hỏi đúng tài liệu này: câu trả lời có ngữ cảnh → ghi progress theo môn.
    client.post(
        "/api/chat",
        json={"message": "phan phoi xac suat la gi", "document_id": document_id},
        headers=auth_header(token),
    )
    body = client.get("/api/progress", headers=auth_header(token)).json()
    assert body["suggestion"] is None
    assert any(
        subject["subject"] == "Xác suất thống kê" and subject["questions_asked"] >= 1
        for subject in body["subjects"]
    )
